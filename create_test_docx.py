import docx

doc = docx.Document()
doc.add_heading("APOLLO DIAGNOSTICS - REPORT", 0)

doc.add_paragraph("Patient Name: {{patient_name}}")
doc.add_paragraph("Age: {{age}}")
doc.add_paragraph("Sex: {{sex}}")
doc.add_paragraph("Referral Doctor: {{ref_doctor}}")
doc.add_paragraph("Study: {{study}}")

doc.add_heading("Findings:", level=1)
doc.add_paragraph("{{findings}}")

doc.add_heading("Impression:", level=1)
doc.add_paragraph("{{impression}}")

doc.save("test_template.docx")
print("test_template.docx generated successfully!")
